"""Build the owner's handover package from the agents' final files.

Usage: python3 trading/tools/build_handover.py
Outputs: trading/HANDOVER.md and trading/TRADING-RESULTS.xlsx
Source of truth = each market's trade list (written by Agent 2) and FINAL-STRATEGY.md. Weekly and yearly tables
are rebuilt the same way for every market, so all four can be compared side by side.
"""
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

T = Path(__file__).resolve().parents[1]

MARKETS = [
    # key, folder, strategy name, grade, trade files, R column, period column, which period labels are "unseen"
    ("NASDAQ", "NASDAQ", "Fade the big opening gap (v1)", "B: Promising",
     ["backtest/hidden_H/trades_v1.csv", "backtest/trades_v1_withC.csv"], "r", "period", {"H", "C", "A", "B"}),
    ("GOLD", "XAUUSD", "Follow London's break of the Asia box (v3)", "B: Promising (bottom of B)",
     ["backtest/trades_v3.csv"], "R", "src", {"hidden"}),
    ("GBPUSD", "GBPUSD", "Quiet-Asia breakout (v3)", "C: Experimental (demo only)",
     ["backtest/trades_v3.csv"], "R", "period", {"hidden"}),
    ("EURUSD", "EURUSD", "Euro-only move fade (v3)", "C: Experimental (demo only)",
     ["backtest/trades_v3.csv"], "R", "period", {"H"}),
]
# Stretches that were never simulated for a market (shown as "not tested" in the weekly sheet)
NOT_TESTED = {"NASDAQ": [("2019-01-01", "2021-10-01")]}


def load(folder, files, rcol, pcol, unseen):
    df = pd.concat([pd.read_csv(T / folder / f) for f in files], ignore_index=True)
    df["date"] = pd.to_datetime(df["date"])
    df = df.drop_duplicates("date").sort_values("date")
    df["R"] = df[rcol].astype(float)
    df["data"] = df[pcol].astype(str).map(lambda p: "Unseen test" if p in unseen else "Study data")
    return df


def weekly(df, key=""):
    start = (df.date.min() - pd.Timedelta(days=df.date.min().weekday())).normalize()
    weeks = pd.date_range(start, df.date.max(), freq="7D")
    rows, run = [], 0.0
    for w in weeks:
        g = df[(df.date >= w) & (df.date < w + pd.Timedelta(days=7))]
        run += g.R.sum()
        data = "/".join(sorted(set(g.data))) if len(g) else ""
        if any(pd.Timestamp(a) <= w < pd.Timestamp(b) for a, b in NOT_TESTED.get(key, [])):
            data = "not tested"
        rows.append({"Week (Mon - Fri)": f"{w:%d %b %Y} - {w + pd.Timedelta(days=4):%d %b %Y}", "Data": data,
                     "Trades": len(g), "Wins": int((g.R > 0).sum()), "Losses": int((g.R <= 0).sum()),
                     "Win rate": f"{(g.R > 0).mean() * 100:.0f}%" if len(g) else "-",
                     "Result (R)": round(g.R.sum(), 2), "Running total (R)": round(run, 2)})
    return pd.DataFrame(rows)


def stats(x):
    if x.empty:
        return {"Trades": 0}
    eq = x.R.cumsum()
    dd = (eq - eq.cummax()).min()
    streak = longest = 0
    for r in x.R:
        streak = streak + 1 if r <= 0 else 0
        longest = max(longest, streak)
    wins, losses = x.R[x.R > 0].sum(), -x.R[x.R <= 0].sum()
    return {"Trades": len(x), "Win rate": f"{(x.R > 0).mean() * 100:.0f}%", "Total (R)": round(x.R.sum(), 1),
            "Per trade (R)": round(x.R.mean(), 3), "Profit factor": round(wins / losses, 2) if losses else None,
            "Biggest drop (R)": round(dd, 1), "Longest losing run": longest}


def yearly(df):
    rows = [{"Year": y, "Data": g.data.iloc[0] if g.data.nunique() == 1 else "mixed (unseen + study)", **stats(g)}
            for y, g in df.groupby(df.date.dt.year)]
    for label in ("Unseen test", "Study data"):
        rows.append({"Year": f"ALL {label.upper()}", "Data": label, **stats(df[df.data == label])})
    rows.append({"Year": "ALL YEARS", "Data": "", **stats(df)})
    return pd.DataFrame(rows)


def strip_weekly_section(md):
    """Remove the long week-by-week table from a FINAL-STRATEGY (it is in the spreadsheet)."""
    out, skip = [], False
    for line in md.splitlines():
        if re.match(r"^###? .*[Ww]eek-by-week", line):
            skip = True
            out.append(line)
            out.append("")
            out.append("*The full week-by-week table is in `TRADING-RESULTS.xlsx` (one sheet per market).*")
            out.append("")
            continue
        if skip and re.match(r"^#{1,3} ", line):
            skip = False
        if not skip:
            out.append(line)
    return "\n".join(out)


def style(ws, widths=None):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), (widths or {}).get(i, 45))


def main():
    data = {k: load(f, files, r, p, u) for k, f, _, _, files, r, p, u in MARKETS}
    summary = []
    for k, folder, name, grade, *_ in MARKETS:
        df = data[k]
        un, st = stats(df[df.data == "Unseen test"]), stats(df[df.data == "Study data"])
        summary.append({"Market": k, "Strategy": name, "Grade": grade,
                        "Years tested": f"{df.date.min():%b %Y} - {df.date.max():%b %Y}",
                        "Unseen trades": un["Trades"], "Unseen total (R)": un.get("Total (R)"),
                        "Unseen per trade (R)": un.get("Per trade (R)"),
                        "Study trades": st["Trades"], "Study total (R)": st.get("Total (R)"),
                        "Study per trade (R)": st.get("Per trade (R)"),
                        "Longest losing run (all)": stats(df)["Longest losing run"],
                        "Strategy document": f"trading/{folder}/FINAL-STRATEGY.md"})
    summary = pd.DataFrame(summary)

    xlsx = T / "TRADING-RESULTS.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        notes = pd.DataFrame({"Read me": [
            "Day-trading research results, prepared by the team (Agent 1 research, Agent 2 testing). 1 R = the 1% of the account risked per trade.",
            "'Unseen test' = years nobody looked at when the rules were written (the real judge). 'Study data' = years the rules were built on.",
            "Grades: A = proven on unseen data; B = made money on unseen data but the edge is small or uneven; C = lost on unseen data (practice only).",
            "Costs included: Nasdaq 2.0 points, gold 0.40 USD, GBPUSD 1.5 pips, EURUSD 1.2 pips per trade. Your broker's real costs may differ.",
            "Prices: HistData.com 1-minute data, clock checked. Gaps (e.g. Feb-Jul 2023) were skipped, so some weeks show no trades.",
            "Sheets: Summary; per market a Weekly sheet (every week), a Yearly sheet and a Trades sheet (every trade); Paper = live practice journal.",
            "Past results do not guarantee future results. Trading carries a real risk of loss. This is not financial advice."]})
        notes.to_excel(xw, sheet_name="Read me", index=False)
        style(xw.sheets["Read me"], {1: 140})
        summary.to_excel(xw, sheet_name="Summary", index=False)
        style(xw.sheets["Summary"])
        for k, *_ in MARKETS:
            df = data[k]
            weekly(df, k).to_excel(xw, sheet_name=f"{k} weekly", index=False)
            style(xw.sheets[f"{k} weekly"])
            yearly(df).to_excel(xw, sheet_name=f"{k} yearly", index=False)
            style(xw.sheets[f"{k} yearly"])
            tr = df.drop(columns=["data"]).copy()
            tr.insert(1, "Data", df.data)
            tr["date"] = tr.date.dt.strftime("%Y-%m-%d")
            tr.to_excel(xw, sheet_name=f"{k} trades", index=False)
            style(xw.sheets[f"{k} trades"], {i: 18 for i in range(1, 60)})
        pj = T / "paper" / "journal.csv"
        if pj.exists():
            pd.read_csv(pj).to_excel(xw, sheet_name="Paper (live practice)", index=False)
            style(xw.sheets["Paper (live practice)"])
    print("wrote", xlsx)

    # ---------------------------------------------------------------- HANDOVER.md
    L = ["# Day-trading project: handover", "",
         f"Prepared by the team leader, {pd.Timestamp.now(tz='Asia/Qatar'):%d %b %Y} (Doha). "
         "Research by Agent 1 (Trading Research Lead), testing and grading by Agent 2 (Backtest & Performance Manager).", "",
         "**What you are getting:** one strategy per market, each with the reasons, the test results and the full A-to-Z rules "
         "(Doha and New York times), graded honestly. All week-by-week results and every single trade are in "
         "**`TRADING-RESULTS.xlsx`**.", "",
         "*Past results do not guarantee future results. Trading carries a real risk of loss. This is not financial advice.*", "",
         "## 1. Results at a glance", "",
         "1 R = the 1% of your account risked on one trade. **Unseen** = years nobody looked at when the rules were written "
         "(the real judge). **Study** = years the rules were built on.", "",
         "| Market | Strategy | Grade | Unseen years | Study years | Longest losing run |", "|---|---|---|---|---|---|"]
    for _, s in summary.iterrows():
        L.append(f"| **{s.Market}** | {s.Strategy} | **{s.Grade}** | {s['Unseen trades']} trades, {s['Unseen total (R)']:+.1f} R "
                 f"({s['Unseen per trade (R)']:+.3f} per trade) | {s['Study trades']} trades, {s['Study total (R)']:+.1f} R "
                 f"({s['Study per trade (R)']:+.3f} per trade) | {s['Longest losing run (all)']} |")
    L += ["", "**Grades:** A = proven on unseen data. B = made money on unseen data, but the edge is small or uneven: "
          "practise it first. C = lost money on unseen data: practice account only, never real money.", "",
          "### Results per year", ""]
    for k, *_ in MARKETS:
        y = yearly(data[k])
        y = y[~y.Year.astype(str).str.startswith("ALL")]
        tag = {"Unseen test": " (unseen)", "Study data": ""}
        L.append(f"**{k}:** " + ", ".join(f"{r.Year} {r['Total (R)']:+.1f}{tag.get(r.Data, ' (part unseen)')}"
                                          for _, r in y.iterrows()))
        L.append("")
    L += ["(Nasdaq was not simulated for 2019 - Sep 2021. Partial years: 2023 has a data gap Feb - Jul; 2026 ends in Sep.)", "",
          "## 2. What to do next", "",
          "1. **Nasdaq and gold (grade B):** check about 20 trades from the `Trades` sheets on your broker's chart, then follow "
          "the live practice journal (below) or practise them yourself on a demo account for at least 3 months. Each strategy's "
          "grade box says exactly what upgrades it to A and when to stop using it.",
          "2. **GBPUSD and EURUSD (grade C):** practice account only. They lost money on unseen years.",
          "3. **Tell us your broker and its costs.** All four results depend on real trading costs; gold and the currency pairs "
          "especially.",
          "4. **Going live is your decision**, after your own backtest and practice.", "",
          "## 3. Live practice trading (running now)", "",
          "Every weekday after the US close (23:20 Doha in US summer, 00:20 in winter) a routine reads that day's live 1-minute "
          "prices for the Nasdaq 100 and gold and records what each grade-B strategy would have done. **No orders are placed.** "
          "Every Sunday morning you get a short report. Journal: `trading/paper/journal.csv`; details: `trading/paper/README.md`.", ""]
    pj = T / "paper" / "journal.csv"
    if pj.exists():
        j = pd.read_csv(pj).fillna("")
        L += ["| Date | Market | What happened | Result |", "|---|---|---|---|"]
        for _, r in j.iterrows():
            what = (f"{r.direction} at {r.entry_time_doha} Doha, closed {r.exit_time_doha} ({r.exit_type})"
                    if r.status == "trade" else f"no trade: {r.reason}")
            L.append(f"| {r.date_ny} | {r.market} | {what} | {f'{float(r.R):+.2f} R' if r.status == 'trade' else '-'} |")
        L.append("")
    L += ["## 4. How we tested (so you can trust the numbers)", "",
          "- **Data:** 1-minute prices, Jan 2016 - Sep 2026 per market (about 11 years). Every file's clock was checked against "
          "US news releases and the 09:30 New York stock open, including the weeks when US and UK clocks differ. Days with "
          "missing prices (mostly Feb - Jul 2023) were skipped.",
          "- **Two separate agents:** Agent 1 built each strategy; Agent 2 rewrote the test from the written rules alone. "
          "Every strategy's trades matched between the two.",
          "- **Unseen years:** each strategy was judged on years it had never seen. Many ideas that looked good on the study years "
          "failed there: over 290 ideas and variations were tried across the four markets. These four are the best found.",
          "- **Costs** were charged on every trade, and every strategy was also tested at double cost.", "",
          "## 5. The four strategies in full", ""]
    for k, folder, *_ in MARKETS:
        md = (T / folder / "FINAL-STRATEGY.md").read_text()
        md = strip_weekly_section(md)
        L += ["---", "", f"<a id='{k.lower()}'></a>", "", re.sub(r"^(#{1,5}) ", lambda m: "#" + m.group(1) + " ", md, flags=re.M), ""]
    (T / "HANDOVER.md").write_text("\n".join(L) + "\n")
    print("wrote", T / "HANDOVER.md")


if __name__ == "__main__":
    main()
